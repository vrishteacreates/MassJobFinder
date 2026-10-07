import streamlit as st
import pandas as pd
from io import BytesIO
from urllib.parse import quote_plus

from job_scraper import search_jobs
from remotive_scraper import search_remotive_jobs


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Mass Job Finder",
    page_icon="💼",
    layout="wide"
)


# ============================================================
# HEADER
# ============================================================

st.title("💼 Mass Job Finder")

st.caption(
    "Find jobs from multiple sources in one place."
)


# ============================================================
# SEARCH INPUTS
# ============================================================

st.subheader("🔎 Find Your Next Job")

col1, col2 = st.columns(2)

with col1:
    keyword = st.text_input(
        "Job Keyword",
        placeholder="Example: Python Developer"
    )

with col2:
    location = st.text_input(
        "Location",
        value="Mumbai",
        placeholder="Example: Mumbai"
    )


# ============================================================
# FILTERS
# ============================================================

col1, col2, col3 = st.columns(3)

with col1:
    experience = st.selectbox(
        "Experience",
        [
            "Any Experience",
            "Fresher",
            "0–1 years",
            "1–3 years",
            "3–5 years",
            "5+ years"
        ]
    )

with col2:
    job_type = st.selectbox(
        "Job Type",
        [
            "Any Type",
            "Full-time",
            "Part-time",
            "Internship",
            "Contract"
        ]
    )

with col3:
    work_mode = st.selectbox(
        "Work Mode",
        [
            "Any",
            "Remote / Work From Home",
            "On-site"
        ]
    )


col1, col2, col3 = st.columns(3)

with col1:
    minimum_salary = st.selectbox(
        "Minimum Salary",
        [
            "Any Salary",
            "₹2 LPA+",
            "₹3 LPA+",
            "₹5 LPA+",
            "₹8 LPA+",
            "₹10 LPA+"
        ]
    )

with col2:
    sort_option = st.selectbox(
        "Sort Jobs",
        [
            "Newest First",
            "Oldest First"
        ]
    )

with col3:
    source_filter = st.selectbox(
        "Job Source",
        [
            "All Sources",
            "Adzuna",
            "Remotive"
        ]
    )


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_work_mode(job):

    text = (
        str(job.get("title", "")) + " " +
        str(job.get("description", "")) + " " +
        str(job.get("location", "")) + " " +
        str(job.get("work_mode", ""))
    ).lower()

    remote_words = [
        "remote",
        "work from home",
        "wfh",
        "work-from-home",
        "worldwide",
        "anywhere"
    ]

    for word in remote_words:

        if word in text:
            return "Remote"

    return "On-site"


def matches_experience(
    job,
    selected_experience
):

    if selected_experience == "Any Experience":
        return True

    text = (
        str(job.get("title", "")) + " " +
        str(job.get("description", ""))
    ).lower()

    experience_keywords = {

        "Fresher": [
            "fresher",
            "freshers",
            "entry level",
            "entry-level",
            "graduate",
            "graduates",
            "no experience"
        ],

        "0–1 years": [
            "0-1 year",
            "0–1 year",
            "0 year",
            "1 year",
            "fresher",
            "entry level",
            "entry-level"
        ],

        "1–3 years": [
            "1-3 year",
            "1–3 year",
            "1 year",
            "2 year",
            "3 year",
            "1 years",
            "2 years",
            "3 years"
        ],

        "3–5 years": [
            "3-5 year",
            "3–5 year",
            "3 years",
            "4 years",
            "5 years"
        ],

        "5+ years": [
            "5+ year",
            "5 years",
            "6 years",
            "7 years",
            "8 years",
            "9 years",
            "10 years",
            "senior",
            "lead",
            "manager"
        ]
    }

    keywords = experience_keywords.get(
        selected_experience,
        []
    )

    return any(
        word in text
        for word in keywords
    )


def matches_job_type(
    job,
    selected_type
):

    if selected_type == "Any Type":
        return True

    text = (
        str(job.get("title", "")) + " " +
        str(job.get("description", "")) + " " +
        str(job.get("job_type", ""))
    ).lower()

    if selected_type == "Full-time":

        return (
            "full-time" in text
            or "full time" in text
        )

    if selected_type == "Part-time":

        return (
            "part-time" in text
            or "part time" in text
        )

    if selected_type == "Internship":

        return "intern" in text

    if selected_type == "Contract":

        return "contract" in text

    return True


def salary_threshold(value):

    values = {

        "Any Salary": 0,

        "₹2 LPA+": 200000,

        "₹3 LPA+": 300000,

        "₹5 LPA+": 500000,

        "₹8 LPA+": 800000,

        "₹10 LPA+": 1000000
    }

    return values.get(
        value,
        0
    )


def get_date(job):

    value = job.get(
        "created",
        ""
    )

    try:

        return pd.to_datetime(
            value,
            utc=True
        )

    except Exception:

        return pd.Timestamp(
            "1970-01-01",
            tz="UTC"
        )


def remove_duplicates(jobs):

    seen = set()

    result = []

    for job in jobs:

        url = str(
            job.get(
                "url",
                ""
            )
        ).strip()

        title = str(
            job.get(
                "title",
                ""
            )
        ).strip().lower()

        company = str(
            job.get(
                "company",
                ""
            )
        ).strip().lower()

        if url:

            key = url

        else:

            key = (
                title +
                "|" +
                company
            )

        if key not in seen:

            seen.add(key)

            result.append(job)

    return result


# ============================================================
# MAJOR PLATFORM LINKS
# ============================================================

def create_platform_links(
    keyword,
    location
):

    keyword_encoded = quote_plus(
        keyword.strip()
    )

    location_encoded = quote_plus(
        location.strip()
    )

    full_query = quote_plus(
        f"{keyword} {location}"
    )

    return {

        "LinkedIn Jobs":
            (
                "https://www.linkedin.com/jobs/search/"
                f"?keywords={keyword_encoded}"
                f"&location={location_encoded}"
            ),

        "Google Jobs":
            (
                "https://www.google.com/search"
                f"?q={full_query}+jobs"
            ),

        "Naukri":
            (
                "https://www.google.com/search"
                f"?q=site%3Anaukri.com+"
                f"{full_query}+jobs"
            ),

        "Internshala":
            (
                "https://www.google.com/search"
                f"?q=site%3Ainternshala.com+"
                f"{full_query}+jobs"
            ),

        "Indeed":
            (
                "https://in.indeed.com/jobs"
                f"?q={keyword_encoded}"
                f"&l={location_encoded}"
            )
    }


# ============================================================
# MAJOR PLATFORMS
# ============================================================

st.divider()

st.subheader("🌐 Search Major Job Platforms")

if keyword.strip():

    st.write(
        f"Search **{keyword}** jobs in **{location}**."
    )

    platforms = create_platform_links(
        keyword,
        location
    )

    col1, col2, col3, col4, col5 = st.columns(5)

    platform_columns = [
        col1,
        col2,
        col3,
        col4,
        col5
    ]

    for index, (
        platform,
        url
    ) in enumerate(
        platforms.items()
    ):

        with platform_columns[index]:

            st.link_button(
                f"🔎 {platform}",
                url,
                use_container_width=True
            )

else:

    st.info(
        "Enter a job keyword above to activate "
        "major-platform searches."
    )


# ============================================================
# SEARCH BUTTON
# ============================================================

st.divider()

search_button = st.button(
    "🚀 Search Available Jobs",
    type="primary",
    use_container_width=True
)


# ============================================================
# SEARCH ADZUNA + REMOTIVE
# ============================================================

if search_button:

    if not keyword.strip():

        st.warning(
            "Please enter a job keyword."
        )

    else:

        all_jobs = []

        # ====================================================
        # ADZUNA
        # ====================================================

        if source_filter in [
            "All Sources",
            "Adzuna"
        ]:

            with st.spinner(
                "Searching Adzuna..."
            ):

                try:

                    adzuna_jobs = search_jobs(
                        keyword=keyword,
                        location=location,
                        experience="Any Experience",
                        job_type="Any Type",
                        results_per_page=50
                    )

                    for job in adzuna_jobs:

                        job["source"] = "Adzuna"

                        job["work_mode"] = (
                            get_work_mode(job)
                        )

                    all_jobs.extend(
                        adzuna_jobs
                    )

                except Exception as error:

                    st.error(
                        f"Adzuna error: {error}"
                    )

        # ====================================================
        # REMOTIVE
        # ====================================================

        if source_filter in [
            "All Sources",
            "Remotive"
        ]:

            with st.spinner(
                "Searching Remotive..."
            ):

                try:

                    remotive_jobs = (
                        search_remotive_jobs(
                            keyword=keyword,
                            limit=50
                        )
                    )

                    for job in remotive_jobs:

                        job["source"] = "Remotive"

                        job["work_mode"] = "Remote"

                    all_jobs.extend(
                        remotive_jobs
                    )

                except Exception as error:

                    st.error(
                        f"Remotive error: {error}"
                    )

        # ====================================================
        # REMOVE DUPLICATES
        # ====================================================

        all_jobs = remove_duplicates(
            all_jobs
        )

        # ====================================================
        # EXPERIENCE FILTER
        # ====================================================

        all_jobs = [

            job

            for job in all_jobs

            if matches_experience(
                job,
                experience
            )
        ]

        # ====================================================
        # JOB TYPE FILTER
        # ====================================================

        all_jobs = [

            job

            for job in all_jobs

            if matches_job_type(
                job,
                job_type
            )
        ]

        # ====================================================
        # WORK MODE
        # ====================================================

        if work_mode == (
            "Remote / Work From Home"
        ):

            all_jobs = [

                job

                for job in all_jobs

                if get_work_mode(job)
                == "Remote"
            ]

        elif work_mode == "On-site":

            all_jobs = [

                job

                for job in all_jobs

                if get_work_mode(job)
                == "On-site"
            ]

        # ====================================================
        # SALARY
        # ====================================================

        minimum_salary_value = (
            salary_threshold(
                minimum_salary
            )
        )

        if minimum_salary_value > 0:

            filtered_jobs = []

            for job in all_jobs:

                salary_min = job.get(
                    "salary_min"
                )

                try:

                    if salary_min is not None:

                        if float(
                            salary_min
                        ) >= minimum_salary_value:

                            filtered_jobs.append(
                                job
                            )

                except Exception:

                    pass

            all_jobs = filtered_jobs

        # ====================================================
        # SORT
        # ====================================================

        all_jobs.sort(
            key=get_date,
            reverse=(
                sort_option == "Newest First"
            )
        )

        # ====================================================
        # SAVE
        # ====================================================

        st.session_state["jobs"] = all_jobs


# ============================================================
# DISPLAY RESULTS
# ============================================================

if "jobs" in st.session_state:

    jobs = st.session_state["jobs"]

    st.divider()

    st.subheader(
        "📋 Available Job Results"
    )

    if not jobs:

        st.warning(
            "No jobs found from the selected sources."
        )

    else:

        # ====================================================
        # METRICS
        # ====================================================

        total_jobs = len(jobs)

        adzuna_count = sum(
            1
            for job in jobs
            if job.get("source") == "Adzuna"
        )

        remotive_count = sum(
            1
            for job in jobs
            if job.get("source") == "Remotive"
        )

        remote_count = sum(
            1
            for job in jobs
            if get_work_mode(job) == "Remote"
        )

        c1, c2, c3, c4 = st.columns(4)

        with c1:
            st.metric(
                "Total Jobs",
                total_jobs
            )

        with c2:
            st.metric(
                "Adzuna",
                adzuna_count
            )

        with c3:
            st.metric(
                "Remotive",
                remotive_count
            )

        with c4:
            st.metric(
                "Remote",
                remote_count
            )

        # ====================================================
        # DATAFRAME
        # ====================================================

        table_data = []

        for job in jobs:

            table_data.append({

                "Job Title":
                    job.get(
                        "title",
                        "Not provided"
                    ),

                "Company":
                    job.get(
                        "company",
                        "Not provided"
                    ),

                "Location":
                    job.get(
                        "location",
                        "Not provided"
                    ),

                "Salary":
                    job.get(
                        "salary",
                        "Not specified"
                    ),

                "Job Type":
                    job.get(
                        "job_type",
                        "Not specified"
                    ),

                "Source":
                    job.get(
                        "source",
                        "Unknown"
                    ),

                "Work Mode":
                    get_work_mode(job),

                "Posted":
                    job.get(
                        "created",
                        ""
                    ),

                "Job Link":
                    job.get(
                        "url",
                        ""
                    )
            })

        df = pd.DataFrame(
            table_data
        )

        st.dataframe(

            df,

            use_container_width=True,

            hide_index=True,

            column_config={

                "Job Link":
                    st.column_config.LinkColumn(
                        "Job Link",
                        display_text="🔗 Open Job"
                    )
            }
        )

        # ====================================================
        # CSV
        # ====================================================

        csv_data = df.to_csv(
            index=False
        ).encode("utf-8")

        st.download_button(

            "⬇️ Download CSV",

            data=csv_data,

            file_name=(
                "mass_job_finder_results.csv"
            ),

            mime="text/csv"
        )

        # ====================================================
        # EXCEL
        # ====================================================

        excel_buffer = BytesIO()

        with pd.ExcelWriter(
            excel_buffer,
            engine="openpyxl"
        ) as writer:

            df.to_excel(
                writer,
                index=False,
                sheet_name="Jobs"
            )

        excel_buffer.seek(0)

        from openpyxl import load_workbook

        workbook = load_workbook(
            excel_buffer
        )

        worksheet = workbook["Jobs"]

        job_link_column = None

        for cell in worksheet[1]:

            if cell.value == "Job Link":

                job_link_column = cell.column

                break

        if job_link_column:

            for row in range(
                2,
                worksheet.max_row + 1
            ):

                cell = worksheet.cell(
                    row=row,
                    column=job_link_column
                )

                url = cell.value

                if url:

                    cell.value = (
                        f'=HYPERLINK('
                        f'"{url}",'
                        f'"Open Job")'
                    )

                    cell.hyperlink = url

                    cell.style = "Hyperlink"

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:

                try:

                    length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        length
                    )

                except Exception:

                    pass

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                50
            )

        final_excel = BytesIO()

        workbook.save(
            final_excel
        )

        final_excel.seek(0)

        st.download_button(

            "📊 Download Excel",

            data=final_excel,

            file_name=(
                "mass_job_finder_results.xlsx"
            ),

            mime=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        # ====================================================
        # JOB DETAILS
        # ====================================================

        st.subheader(
            "📌 Job Details"
        )

        for index, job in enumerate(jobs):

            title = job.get(
                "title",
                "Untitled Job"
            )

            company = job.get(
                "company",
                "Company not provided"
            )

            source = job.get(
                "source",
                "Unknown"
            )

            with st.expander(
                f"{index + 1}. "
                f"{title} — "
                f"{company} "
                f"({source})"
            ):

                st.write(
                    f"**Company:** {company}"
                )

                st.write(
                    f"**Location:** "
                    f"{job.get('location', 'Not provided')}"
                )

                st.write(
                    f"**Salary:** "
                    f"{job.get('salary', 'Not specified')}"
                )

                st.write(
                    f"**Job Type:** "
                    f"{job.get('job_type', 'Not specified')}"
                )

                st.write(
                    f"**Work Mode:** "
                    f"{get_work_mode(job)}"
                )

                st.write(
                    f"**Source:** {source}"
                )

                st.write(
                    job.get(
                        "description",
                        "No description available."
                    )
                )

                url = job.get(
                    "url",
                    ""
                )

                if url:

                    st.link_button(
                        "🔗 Apply / View Job",
                        url
                    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "Mass Job Finder combines available job data "
    "with direct searches on major job platforms."
)